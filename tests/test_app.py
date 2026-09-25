#!/usr/bin/env python3
"""Flask route tests for the web app (no network calls).

Run from the repo root:
    python -m unittest discover -s tests
"""

import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import registrar_io
import app as app_module
from app import app


def _registrar_xlsx_bytes():
    """Build a tiny valid registrar .xlsx in memory."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Export'
    ws.append(registrar_io.REGISTRAR_COLUMNS)
    from datetime import datetime
    ws.append([None, 'GEOG', '1001', '10', 'Intro to Human Geography', None,
               'G10715190', 'Chacko', 'Elizabeth', 3, 120, 117, 40,
               datetime(2027, 1, 11), datetime(2027, 4, 26), 'TR', '1420', '1535', None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestRoutes(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_index_ok(self):
        r = self.client.get('/')
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'Scrape a subject', r.data)
        self.assertIn(b'Upload registrar', r.data)

    def test_health(self):
        r = self.client.get('/health')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data, b'ok')

    def test_scrape_bad_subject_400(self):
        # Invalid subject -> build_gwu_url raises -> 400, no network hit.
        r = self.client.post('/scrape', data={'year': '2026', 'semester': '01', 'subject': '1'})
        self.assertEqual(r.status_code, 400)

    def test_scrape_bad_year_400(self):
        r = self.client.post('/scrape', data={'year': '26', 'semester': '01', 'subject': 'GEOG'})
        self.assertEqual(r.status_code, 400)

    def test_upload_non_xlsx_400(self):
        r = self.client.post('/upload', data={
            'file': (io.BytesIO(b'not a spreadsheet'), 'notes.txt')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)

    def test_upload_valid_renders_calendar(self):
        r = self.client.post('/upload', data={
            'file': (io.BytesIO(_registrar_xlsx_bytes()), 'sched.xlsx')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'text/html', r.content_type.encode())
        # The web-export button should be present in web mode.
        self.assertIn(b'exportToRegistrarServer', r.data)
        self.assertIn(b'Chacko', r.data)

    def test_export_xlsx_ok(self):
        courses = [{
            'subject': 'GEOG', 'course_num': '1001', 'title': 'Intro',
            'instructor': 'Chacko, E', 'credits': '3.00', 'days': 'TR',
            'time': {'start': '02:20PM', 'end': '03:35PM', 'raw': '02:20PM - 03:35PM'},
            'dates': '01/11/27 - 04/26/27', 'course_number': 'GEOG 1001',
        }]
        r = self.client.post('/export/xlsx', json=courses)
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r.content_type)
        self.assertTrue(r.data[:2] == b'PK')  # .xlsx is a zip

    def test_export_xlsx_bad_payload_400(self):
        self.assertEqual(self.client.post('/export/xlsx', json={'not': 'a list'}).status_code, 400)
        self.assertEqual(self.client.post('/export/xlsx', json=['not a dict']).status_code, 400)

    def test_security_header(self):
        r = self.client.get('/')
        self.assertEqual(r.headers.get('X-Content-Type-Options'), 'nosniff')


_COURSE = {'crn': '46234', 'subject': 'GEOG', 'course_num': '1001', 'section': '10',
           'title': 'Intro to Human Geography', 'credits': '3.00', 'instructor': 'Chacko, E',
           'days': 'TR', 'time': {'start': '02:20PM', 'end': '03:35PM',
                                  'raw': '02:20PM - 03:35PM'},
           'dates': '01/12/26 - 04/27/26', 'building': '1957 E', 'room': 'B12',
           'status': 'OPEN', 'course_number': 'GEOG 1001'}


def _xlsx_with_rows(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(registrar_io.REGISTRAR_COLUMNS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestScrapeRoute(unittest.TestCase):
    """/scrape with the network replaced by a stub scraper."""

    def setUp(self):
        self.client = app.test_client()

    def _post(self, **form):
        data = {'year': '2026', 'semester': '03', 'subject': 'geog'}
        data.update(form)
        return self.client.post('/scrape', data=data)

    def test_success_renders_calendar_for_validated_url(self):
        with mock.patch.object(app_module, 'CourseScraper') as cls:
            cls.return_value.scrape.return_value = [_COURSE]
            r = self._post()
        self.assertEqual(r.status_code, 200)
        cls.assert_called_once_with(
            url='https://my.gwu.edu/mod/pws/courses.cfm?campId=1&termId=202603&subjId=GEOG')
        self.assertIn(b'GWU Courses - Fall 2026', r.data)
        self.assertIn(b'Chacko, E', r.data)
        self.assertIn(b'exportToRegistrarServer', r.data)

    def test_scraper_crash_is_502(self):
        with mock.patch.object(app_module, 'CourseScraper') as cls:
            cls.return_value.scrape.side_effect = RuntimeError('boom')
            r = self._post()
        self.assertEqual(r.status_code, 502)
        self.assertIn(b"reach the GWU schedule site", r.data)

    def test_no_courses_is_404(self):
        with mock.patch.object(app_module, 'CourseScraper') as cls:
            cls.return_value.scrape.return_value = []
            r = self._post()
        self.assertEqual(r.status_code, 404)
        self.assertIn(b'No courses found', r.data)

    def test_bad_semester_400_without_scraping(self):
        with mock.patch.object(app_module, 'CourseScraper') as cls:
            r = self._post(semester='08')
        self.assertEqual(r.status_code, 400)
        cls.assert_not_called()

    def test_injection_in_subject_rejected(self):
        with mock.patch.object(app_module, 'CourseScraper') as cls:
            r = self._post(subject='GEOG&campId=2')
        self.assertEqual(r.status_code, 400)
        cls.assert_not_called()


class TestUploadRoute(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def _upload(self, content, name='sched.xlsx'):
        return self.client.post('/upload', data={'file': (io.BytesIO(content), name)},
                                content_type='multipart/form-data')

    def test_missing_file_400(self):
        r = self.client.post('/upload', data={}, content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)
        self.assertIn(b'No file was uploaded', r.data)

    def test_uppercase_extension_accepted(self):
        r = self._upload(_registrar_xlsx_bytes(), name='SCHED.XLSX')
        self.assertEqual(r.status_code, 200)

    def test_corrupt_xlsx_400(self):
        r = self._upload(b'PK\x03\x04 not really a workbook')
        self.assertEqual(r.status_code, 400)
        self.assertIn(b"read that file", r.data)

    def test_wrong_headers_400(self):
        wb = openpyxl.Workbook()
        wb.active.append(['Name', 'Email'])
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(self._upload(buf.getvalue()).status_code, 400)

    def test_header_only_400(self):
        r = self._upload(_xlsx_with_rows([]))
        self.assertEqual(r.status_code, 400)
        self.assertIn(b'No course rows', r.data)

    def test_tba_rows_left_off_calendar(self):
        from datetime import datetime
        rows = [
            [None, 'GEOG', '1001', '10', 'Timed', None, None, 'Chacko', 'E', 3, 10, 1, 0,
             datetime(2027, 1, 11), datetime(2027, 4, 26), 'TR', '1420', '1535', None],
            [None, 'GEOG', '3998', '10', 'Arranged', None, None, 'Mann', 'M', 3, 10, 1, 0,
             datetime(2027, 1, 11), datetime(2027, 4, 26), None, '####', '####', None],
        ]
        r = self._upload(_xlsx_with_rows(rows))
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'Timed', r.data)
        self.assertNotIn(b'Arranged', r.data)

    def test_oversize_upload_413(self):
        big = b'0' * (app.config['MAX_CONTENT_LENGTH'] + 1)
        self.assertEqual(self._upload(big).status_code, 413)


class TestExportRoute(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_too_many_courses_400(self):
        payload = [_COURSE] * (app_module.MAX_EXPORT_COURSES + 1)
        self.assertEqual(self.client.post('/export/xlsx', json=payload).status_code, 400)

    def test_max_courses_ok(self):
        payload = [_COURSE] * app_module.MAX_EXPORT_COURSES
        self.assertEqual(self.client.post('/export/xlsx', json=payload).status_code, 200)

    def test_non_json_body_400(self):
        r = self.client.post('/export/xlsx', data='hello', content_type='text/plain')
        self.assertEqual(r.status_code, 400)

    def test_tba_course_exports_blank_times(self):
        tba = dict(_COURSE, time=None, days='')
        r = self.client.post('/export/xlsx', json=[tba])
        self.assertEqual(r.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertIsNone(row['Begin Time HHMM'])
        self.assertIsNone(row['Weekly Meeting Pattern'])


class TestAnalytics(unittest.TestCase):
    """The GA4 tag is rendered on every server page and can be disabled via env."""

    def setUp(self):
        self.client = app.test_client()

    def test_default_id_on_index(self):
        app.config['GA_MEASUREMENT_ID'] = 'G-BD4W32R0TM'
        r = self.client.get('/')
        self.assertIn(b'googletagmanager.com/gtag/js?id=G-BD4W32R0TM', r.data)
        self.assertIn(b"gtag('config', 'G-BD4W32R0TM')", r.data)

    def test_tag_on_error_page(self):
        app.config['GA_MEASUREMENT_ID'] = 'G-BD4W32R0TM'
        r = self.client.post('/upload', data={
            'file': (io.BytesIO(b'not a spreadsheet'), 'notes.txt')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 400)
        self.assertIn(b'gtag/js?id=G-BD4W32R0TM', r.data)

    def test_tag_on_calendar_page_but_not_in_offline_export(self):
        app.config['GA_MEASUREMENT_ID'] = 'G-BD4W32R0TM'
        r = self.client.post('/upload', data={
            'file': (io.BytesIO(_registrar_xlsx_bytes()), 'sched.xlsx')},
            content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200)
        self.assertIn(b'gtag/js?id=G-BD4W32R0TM', r.data)
        # The embedded offline template (downloaded "Export Schedule") must not
        # carry the tag: decode the base64 blob and check.
        import base64, re
        m = re.search(rb'const EXPORT_TEMPLATE_B64 = "([^"]+)"', r.data)
        self.assertIsNotNone(m)
        offline = base64.b64decode(m.group(1))
        self.assertNotIn(b'googletagmanager', offline)

    def test_empty_id_disables_tag(self):
        app.config['GA_MEASUREMENT_ID'] = ''
        try:
            self.assertNotIn(b'googletagmanager', self.client.get('/').data)
            r = self.client.post('/upload', data={
                'file': (io.BytesIO(_registrar_xlsx_bytes()), 'sched.xlsx')},
                content_type='multipart/form-data')
            self.assertNotIn(b'googletagmanager', r.data)
        finally:
            app.config['GA_MEASUREMENT_ID'] = 'G-BD4W32R0TM'

    def test_env_var_overrides_default(self):
        import importlib, app as app_module
        os.environ['GA_MEASUREMENT_ID'] = 'G-TESTOVERRIDE'
        try:
            importlib.reload(app_module)
            self.assertEqual(app_module.app.config['GA_MEASUREMENT_ID'], 'G-TESTOVERRIDE')
        finally:
            del os.environ['GA_MEASUREMENT_ID']
            importlib.reload(app_module)


if __name__ == '__main__':
    unittest.main()
