#!/usr/bin/env python3
"""Edit-mode tests: a course edited in the calendar must show up correctly in
the registrar .xlsx export, and in the calendar rebuilt from that export.

The browser-side edit logic (openEditCourseModal / saveCourseEdit / ...) is run
for real in Node against a tiny DOM stub, using the JS extracted from the page
the web app actually serves. The resulting ``editedCourses`` is POSTed to
/export/xlsx exactly as the page's export button does. Node-backed tests are
skipped when Node isn't installed.

Run from the repo root:
    python -m unittest discover -s tests
"""

import base64
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
import registrar_io
from app import app

NODE = shutil.which('node')

# Page functions the edit flow needs. Everything else (rendering, filters,
# toasts) is stubbed as a no-op.
JS_FUNCTIONS = ['timeToMinutes', 'openEditCourseModal', 'openAddCourseModal',
                'closeEditModal', 'saveCourseEdit', 'deleteCourse', 'duplicateCourse']
JS_NOOPS = ['updateEditCount', 'populateEditInstructorFilter', 'populateEditRoomDropdown',
            'renderEditCalendar', 'displayEditModeConflicts', 'populateInstructorDatalist',
            'showToast']

# Minimal DOM: elements by id, weekday checkboxes, and a form reset() that
# behaves like the browser's (hidden inputs keep their value).
DOM_STUB = r'''
const HIDDEN_INPUTS = new Set(['editCRN', 'originalCRN']);
const els = {};
function el(id) {
    if (!els[id]) els[id] = { id, value: '', textContent: '', style: {} };
    return els[id];
}
const dayBoxes = ['M', 'T', 'W', 'R', 'F'].map(v => ({ value: v, checked: false }));
el('editCourseForm').reset = function () {
    for (const id of Object.keys(els)) if (!HIDDEN_INPUTS.has(id)) els[id].value = '';
    dayBoxes.forEach(b => b.checked = false);
};
globalThis.document = {
    getElementById: el,
    querySelectorAll(sel) {
        if (sel.includes('editDay')) return sel.includes(':checked') ? dayBoxes.filter(b => b.checked) : dayBoxes;
        return [];
    },
    querySelector(sel) {
        const m = sel.match(/editDay"\]\[value="(\w)"\]/);
        return m ? dayBoxes.find(b => b.value === m[1]) || null : null;
    },
};
globalThis.confirm = () => true;
function setForm(fields, days) {
    for (const [k, v] of Object.entries(fields || {})) el(k).value = v;
    if (days !== undefined) dayBoxes.forEach(b => b.checked = days.includes(b.value));
}
const submit = { preventDefault() {} };
'''

# Each op mimics a user: open the dialog, change fields, click a button.
OPS_RUNNER = r'''
for (const op of OPS) {
    if (op.op === 'add') {
        openAddCourseModal();
    } else {
        openEditCourseModal(editedCourses.find(c => c.crn === op.crn));
    }
    if (op.op === 'edit' || op.op === 'add') {
        setForm(op.fields, op.days);
        saveCourseEdit(submit);
    } else if (op.op === 'delete') {
        deleteCourse();
    } else if (op.op === 'duplicate') {
        duplicateCourse();
    } else if (op.op === 'cancel') {
        closeEditModal();
    }
}
process.stdout.write(JSON.stringify(editedCourses));
'''


def _extract_js_function(html, name):
    """Return the source of ``function name(...) {...}`` from the page."""
    start = html.find(f'function {name}(')
    if start == -1:
        raise AssertionError(f'function {name} not found in calendar page')
    depth = 0
    i = html.index('{', start)
    while True:
        if html[i] == '{':
            depth += 1
        elif html[i] == '}':
            depth -= 1
            if depth == 0:
                return html[start:i + 1]
        i += 1


def _page_courses(html):
    """The course list the calendar page embeds (``const courses = [...];``)."""
    m = re.search(r'const courses = (\[.*?\]);\n', html, re.DOTALL)
    assert m, 'embedded courses not found'
    return json.loads(m.group(1))


def run_edit_ops(html, ops):
    """Run the page's own edit functions on its courses; return editedCourses."""
    parts = [DOM_STUB,
             'const courses = ' + json.dumps(_page_courses(html)) + ';',
             'let editedCourses = JSON.parse(JSON.stringify(courses));',
             'let editedCRNs = new Set(); let editCount = 0; let currentEditCRN = null;']
    parts += [f'function {n}() {{}}' for n in JS_NOOPS]
    parts += [_extract_js_function(html, n) for n in JS_FUNCTIONS]
    parts += ['const OPS = ' + json.dumps(ops) + ';', OPS_RUNNER]
    with tempfile.NamedTemporaryFile('w', suffix='.js', delete=False) as f:
        f.write('\n'.join(parts))
        path = f.name
    try:
        out = subprocess.run([NODE, path], capture_output=True, text=True, timeout=30)
    finally:
        os.unlink(path)
    if out.returncode != 0:
        raise AssertionError(f'node failed:\n{out.stderr}')
    return json.loads(out.stdout)


def _registrar_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Export'
    ws.append(registrar_io.REGISTRAR_COLUMNS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Two courses: a plain one, and a special-topics section with a subtitle.
SOURCE_ROWS = [
    [None, 'GEOG', '1001', '10', 'Intro to Human Geography', None,
     'G10715190', 'Chacko', 'Elizabeth', 3, 120, 117, 40,
     datetime(2027, 1, 11), datetime(2027, 4, 26), 'TR', '1420', '1535', 'keep me'],
    [None, 'GEOG', '3195', '10', 'Special Topics', 'Urban Heat',
     'G20000001', 'Mann', 'Michael', 3, 25, 5, 10,
     datetime(2027, 1, 11), datetime(2027, 4, 26), 'M', '1530', '1800', None],
]


def _export_rows(client, courses):
    """POST courses like the page's export button; return rows keyed by column."""
    r = client.post('/export/xlsx', json=courses)
    assert r.status_code == 200, r.data
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
    return r.data, rows


def _row(rows, course_num, section):
    matches = [r for r in rows if str(r['Course Number']) == course_num
               and str(r['Section']) == section]
    assert len(matches) == 1, f'expected one {course_num}-{section}, got {len(matches)}'
    return matches[0]


def _crn(courses, course_num):
    return next(c['crn'] for c in courses if c['course_num'] == course_num)


@unittest.skipUnless(NODE, 'Node.js not installed')
class TestEditThenExport(unittest.TestCase):
    """Edit in the calendar -> export .xlsx -> re-upload: edits must survive."""

    @classmethod
    def setUpClass(cls):
        cls.client = app.test_client()
        r = cls.client.post('/upload', data={
            'file': (io.BytesIO(_registrar_xlsx(SOURCE_ROWS)), 'sched.xlsx')},
            content_type='multipart/form-data')
        assert r.status_code == 200
        cls.html = r.data.decode('utf-8')
        cls.courses = _page_courses(cls.html)

    def test_edit_every_field_exports_edited_values(self):
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn, 'days': 'MW',
            'fields': {'editSection': '11', 'editTitle': 'Human Geography',
                       'editCredits': '4', 'editInstructor': 'Smith, J',
                       'editStartTime': '09:30AM', 'editEndTime': '10:45AM',
                       'editDates': '01/19/27 - 05/03/27'},
        }])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '1001', '11')
        self.assertEqual(row['Weekly Meeting Pattern'], 'MW')
        self.assertEqual(row['Begin Time HHMM'], '0930')
        self.assertEqual(row['End Time HHMM'], '1045')
        self.assertEqual(row['Course'], 'Human Geography')
        self.assertEqual(row['Credits'], 4)
        self.assertEqual(row['Instructor Last Name'], 'Smith')
        self.assertEqual(row['Instructor First Name'], 'J')
        self.assertEqual(row['Course Start Date'], datetime(2027, 1, 19))
        self.assertEqual(row['Course End Date'], datetime(2027, 5, 3))

    def test_untouched_registrar_fields_survive_an_edit(self):
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn, 'fields': {'editStartTime': '03:00PM',
                                                 'editEndTime': '04:15PM'}}])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '1001', '10')
        self.assertEqual(row['Begin Time HHMM'], '1500')
        # Nothing else about the instructor changed, so the full first name,
        # GWID, enrollment, dates and comment must be exported unchanged.
        self.assertEqual(row['Instructor First Name'], 'Elizabeth')
        self.assertEqual(row['Instructor GWID'], 'G10715190')
        self.assertEqual(row['Max Enrollment'], 120)
        self.assertEqual(row['Prior Enrollment'], 117)
        self.assertEqual(row['Wait Capacity'], 40)
        self.assertEqual(row['Course Start Date'], datetime(2027, 1, 11))
        self.assertEqual(row['Comment'], 'keep me')

    def test_new_instructor_does_not_inherit_old_gwid(self):
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn, 'fields': {'editInstructor': 'Smith, J'}}])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '1001', '10')
        self.assertEqual(row['Instructor Last Name'], 'Smith')
        self.assertIsNone(row['Instructor GWID'])

    def test_edit_special_topics_subtitle(self):
        crn = _crn(self.courses, '3195')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn,
            'fields': {'editTitle': 'Special Topics - Urban Flooding'}}])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '3195', '10')
        self.assertEqual(row['Course'], 'Special Topics')
        self.assertEqual(row['Section Title'], 'Urban Flooding')

    def test_remove_special_topics_subtitle(self):
        crn = _crn(self.courses, '3195')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn, 'fields': {'editTitle': 'Special Topics'}}])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '3195', '10')
        self.assertEqual(row['Course'], 'Special Topics')
        self.assertIsNone(row['Section Title'])

    def test_rename_course_keeps_subtitle(self):
        crn = _crn(self.courses, '3195')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn,
            'fields': {'editTitle': 'Topics in Geography - Urban Heat'}}])
        _, rows = _export_rows(self.client, edited)
        row = _row(rows, '3195', '10')
        self.assertEqual(row['Course'], 'Topics in Geography')
        self.assertEqual(row['Section Title'], 'Urban Heat')

    def test_delete_removes_course_from_export(self):
        crn = _crn(self.courses, '3195')
        edited = run_edit_ops(self.html, [{'op': 'delete', 'crn': crn}])
        _, rows = _export_rows(self.client, edited)
        self.assertEqual([str(r['Course Number']) for r in rows], ['1001'])

    def test_duplicate_adds_new_section_to_export(self):
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [{'op': 'duplicate', 'crn': crn}])
        self.assertEqual(len({c['crn'] for c in edited}), len(edited), 'CRNs must stay unique')
        _, rows = _export_rows(self.client, edited)
        self.assertEqual(_row(rows, '1001', '10')['Begin Time HHMM'], '1420')
        self.assertEqual(_row(rows, '1001', '20')['Begin Time HHMM'], '1420')

    def test_add_new_course_after_editing_another(self):
        # Open an existing course, cancel, then add a new one: the new course
        # must be added, not overwrite the course that was opened first.
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [
            {'op': 'cancel', 'crn': crn},
            {'op': 'add', 'days': 'F',
             'fields': {'editCRNInput': '55555', 'editCourseNum': '2000',
                        'editTitle': 'Map Lab', 'editInstructor': 'Lee, K',
                        'editStartTime': '10:00AM', 'editEndTime': '11:50AM',
                        'editDates': '01/11/27 - 04/26/27'}},
        ])
        self.assertEqual(len(edited), 3)
        _, rows = _export_rows(self.client, edited)
        self.assertEqual(_row(rows, '1001', '10')['Course'], 'Intro to Human Geography')
        new = _row(rows, '2000', '10')
        self.assertEqual(new['Course'], 'Map Lab')
        self.assertEqual(new['Weekly Meeting Pattern'], 'F')
        self.assertEqual(new['Instructor Last Name'], 'Lee')
        self.assertEqual(new['Credits'], 3)

    def test_exported_edits_reload_into_calendar(self):
        # Going forward: re-uploading the exported file shows the edited course.
        crn = _crn(self.courses, '1001')
        edited = run_edit_ops(self.html, [{
            'op': 'edit', 'crn': crn, 'days': 'MW',
            'fields': {'editTitle': 'Human Geography', 'editInstructor': 'Smith, J',
                       'editStartTime': '09:30AM', 'editEndTime': '10:45AM'}}])
        xlsx, _ = _export_rows(self.client, edited)
        r = self.client.post('/upload', data={'file': (io.BytesIO(xlsx), 'edited.xlsx')},
                             content_type='multipart/form-data')
        self.assertEqual(r.status_code, 200)
        reloaded = _page_courses(r.data.decode('utf-8'))
        c = next(c for c in reloaded if c['course_num'] == '1001')
        self.assertEqual(c['title'], 'Human Geography')
        self.assertEqual(c['days'], 'MW')
        self.assertEqual(c['instructor'], 'Smith, J')
        self.assertEqual(c['time']['start'], '09:30AM')
        self.assertEqual(c['time']['end'], '10:45AM')
        self.assertEqual(len(reloaded), 2)

    def test_offline_export_template_carries_edit_functions(self):
        # The downloadable "Export Schedule" page embeds the same page template,
        # so edits keep working in the saved copy.
        m = re.search(r'const EXPORT_TEMPLATE_B64 = "([^"]+)"', self.html)
        offline = base64.b64decode(m.group(1)).decode('utf-8')
        for name in ('saveCourseEdit', 'deleteCourse', 'duplicateCourse'):
            self.assertIn(f'function {name}(', offline)


class TestExportScrapedCourseEdits(unittest.TestCase):
    """Scraped courses (no registrar fields) edited and exported, no Node needed."""

    def test_scraped_course_edit_exports(self):
        edited = {
            'crn': '46234', 'subject': 'GEOG', 'course_num': '1001', 'section': '11',
            'title': 'Human Geography', 'credits': '4.00', 'instructor': 'Smith, J',
            'instructor_last': 'Smith', 'instructor_first': 'J',
            'days': 'MW', 'time': {'start': '09:30AM', 'end': '10:45AM',
                                   'raw': '09:30AM - 10:45AM'},
            'dates': '01/19/27 - 05/03/27', 'building': '1957 E', 'room': 'B12',
            'status': 'OPEN',
        }
        _, rows = _export_rows(app.test_client(), [edited])
        row = _row(rows, '1001', '11')
        self.assertEqual(row['Begin Time HHMM'], '0930')
        self.assertEqual(row['Instructor Last Name'], 'Smith')
        self.assertEqual(row['Course Start Date'], datetime(2027, 1, 19))
        self.assertEqual(row['Credits'], 4)


if __name__ == '__main__':
    unittest.main()
